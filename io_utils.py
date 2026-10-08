from __future__ import annotations
import io
import pandas as pd, numpy as np
from engine import Inputs
import events as ev
import periode as pr

SHEETS = ["Historis", "Dasar", "Pensiun", "Jalur", "Skenario_TPP"]
EVENT_SHEETS = ["Kejadian", "Asumsi_Kejadian"]          # opsional: boleh tidak ada / kosong
EVENT_COLS = {"Kejadian": ev.COLS_K, "Asumsi_Kejadian": ev.COLS_A}
EVENT_NUM = ["tahun_mulai", "nilai"]                     # kolom angka; sisanya teks (tahun_selesai boleh 'akhir')


def coerce_event_table(name: str, df: pd.DataFrame) -> pd.DataFrame:
    """Rapikan tabel kejadian untuk editor UI: kolom baku, angka -> float, teks -> str, baris kosong dibuang."""
    cols = EVENT_COLS[name]
    df = ev._prep(df, cols)
    out = pd.DataFrame(index=range(len(df)))
    for c in cols:
        if c in EVENT_NUM:
            out[c] = pd.to_numeric(df[c], errors="coerce").astype(float)
        else:
            def txt(v):
                if ev._blank(v):
                    return ""
                try:
                    f = float(v)
                    return str(int(f)) if f == int(f) else str(v)
                except (TypeError, ValueError):
                    return str(v).strip()
            out[c] = df[c].map(txt).astype(object)
    keep = ~(out[EVENT_NUM].isna().all(axis=1) & (out.drop(columns=EVENT_NUM) == "").all(axis=1))
    return out[keep].reset_index(drop=True)


def _year_cols(df):
    out = {}
    for c in df.columns:
        try:
            out[c] = int(float(str(c).strip()))
        except ValueError:
            pass
    return out


def tables_from_excel(file) -> dict:
    """Baca template -> dict of DataFrame mentah (untuk ditampilkan/diedit di UI)."""
    xl = pd.ExcelFile(file)
    miss = [s for s in SHEETS if s not in xl.sheet_names]
    if miss:
        raise ValueError(f"Sheet tidak ditemukan: {', '.join(miss)}. Gunakan template resmi aplikasi.")
    t = {s: xl.parse(s) for s in SHEETS}
    for s in EVENT_SHEETS:      # opsional
        t[s] = coerce_event_table(s, xl.parse(s) if s in xl.sheet_names else None)
    if "Setup" in xl.sheet_names:       # opsional: setup yang dipakai saat template dibuat
        t["Setup"] = xl.parse("Setup")
    h = t["Historis"]
    h = h.rename(columns=_year_cols(h))
    t["Historis"] = h
    p = t["Pensiun"]
    t["Pensiun"] = p.rename(columns=_year_cols(p))
    return t


def normalize_tables(t: dict) -> dict:
    """Kolom tahun (teks/angka) -> int, agar tabel hasil edit di UI (kolom string) tetap terbaca."""
    out = {}
    for k, df in t.items():
        df = df.copy()
        if k not in EVENT_SHEETS:
            df.columns = [(_year_cols(df).get(c, c)) for c in df.columns]
        out[k] = df
    return out


def tables_to_ui(t: dict) -> dict:
    out = {}
    for k, df in t.items():
        df = df.copy()
        df.columns = [str(c) for c in df.columns]
        out[k] = df
    return out


def inputs_from_tables(t: dict, setup: "pr.Setup | None" = None, inflasi_dari_data: bool = False) -> Inputs:
    """Tabel -> Inputs. Bila `setup` diberikan: struktur tabel diperiksa terhadap setup, tahun dasar diambil dari setup,
    dan sel Historis yang dibutuhkan model tetapi kosong dilaporkan (bukan diam-diam dianggap 0)."""
    t = normalize_tables(t)
    if setup is not None:
        iss = pr.structure_issues(setup, t)
        if iss:
            raise ValueError("STRUKTUR: " + " ".join(iss) + " Gunakan tombol 'Sesuaikan struktur tabel dengan setup' atau unduh template baru.")
    h = t["Historis"].copy()
    ycols = [c for c in h.columns if isinstance(c, (int, np.integer))]
    h = h.set_index("kode")[sorted(ycols)].apply(pd.to_numeric, errors="coerce")
    h.columns = [int(c) for c in h.columns]
    if setup is not None:
        miss = pr.missing_required(setup, h, need_inflasi=inflasi_dari_data)
        if miss:
            raise ValueError(f"Sheet Historis: {len(miss)} sel kosong yang dibutuhkan model (tahun normal/pemulihan/rujukan/dasar): "
                             + ", ".join(miss[:15]) + (" …" if len(miss) > 15 else ""))
    need = ["inflasi", "pajak", "retribusi", "kekayaan", "lain_pad", "trf_pusat", "trf_antar", "hibah", "pend_lain",
            "peg_total", "barjas", "belanja_hibah", "bansos", "modal", "btt", "transfer", "pemeliharaan"]
    miss = [k for k in need if k not in h.index]
    if miss:
        raise ValueError(f"Kode akun hilang di sheet Historis: {miss}")
    h = h.fillna({k: 0.0 for k in h.index if k != "inflasi"})
    d = t["Dasar"].set_index("kode")["nilai"].to_dict()
    for k in ["gaji_pns", "gaji_pppk", "tpp", "insentif_pajak", "insentif_retribusi", "tunjangan_guru", "asn_tpp_dasar", "silpa", "retribusi_kesehatan", "tahun_dasar"]:
        if k not in d:
            raise ValueError(f"Kode '{k}' hilang di sheet Dasar")
    d = {k: float(v) for k, v in d.items() if not (isinstance(v, str) and not v.strip())}
    if setup is not None:
        d["tahun_dasar"] = setup.tahun_dasar
        blank = [k for k, _ in pr.DASAR_WAJIB if not np.isfinite(d.get(k, np.nan))]
        if blank:
            raise ValueError(f"Sheet Dasar: nilai kosong untuk {blank}.")
    d["tahun_dasar"] = int(d["tahun_dasar"])
    p = t["Pensiun"].copy()
    yc = sorted([c for c in p.columns if isinstance(c, (int, np.integer))])
    tarif = pd.to_numeric(p["gaji_tunjangan_per_tahun"], errors="coerce").fillna(0)
    pen = pd.DataFrame({"orang": [pd.to_numeric(p[y], errors="coerce").fillna(0).sum() for y in yc],
                        "gaji": [float((pd.to_numeric(p[y], errors="coerce").fillna(0) * tarif).sum()) for y in yc]},
                       index=[int(y) for y in yc])
    j = t["Jalur"].copy().set_index("tahun")
    j.index = j.index.astype(int)
    j = j[["asn_penerima_tpp", "gaji_pppk"]].apply(pd.to_numeric, errors="coerce")
    j.columns = ["asn_tpp", "gaji_pppk"]
    if j.isna().any().any():
        raise ValueError("Sheet Jalur: ada sel kosong/non-angka")
    s = t["Skenario_TPP"].copy().set_index("tahun")
    s.index = s.index.astype(int)
    s = s.apply(pd.to_numeric, errors="coerce")
    kj = t.get("Kejadian")
    ak = t.get("Asumsi_Kejadian")
    return Inputs(hist=h, dasar=d, pensiun=pen, jalur=j, tpp_skenario=s,
                  kejadian=ev.empty_k() if kj is None else coerce_event_table("Kejadian", kj),
                  asumsi_kejadian=ev.empty_a() if ak is None else coerce_event_table("Asumsi_Kejadian", ak))


def load_inputs(file):
    return inputs_from_tables(tables_from_excel(file))


# ---------------------------------------------------------------- template Excel adaptif (mengikuti Setup)
def _petunjuk(s: "pr.Setup") -> list:
    j = lambda t: ", ".join(str(y) for y in t) or "-"
    lab = ev.komponen_labels(s.jenis_pemda)
    return [
        "TEMPLATE INPUT — PROYEKSI KEPATUHAN UU HKPD (Pasal 146 & 147)",
        "Template ini dibuat otomatis dari SETUP di sidebar aplikasi. Jangan ubah nama sheet, judul kolom, maupun kolom 'kode'.",
        "Bila setup di aplikasi diubah (tahun awal/dasar/akhir), unduh template baru atau pakai tombol 'Sesuaikan struktur tabel'.",
        "",
        f"SETUP: pemda={s.jenis_pemda} | data historis {s.tahun_awal}–{s.tahun_dasar} | tahun dasar {s.tahun_dasar} | proyeksi {s.tahun_dasar + 1}–{s.tahun_akhir} | target kepatuhan {s.tahun_target}",
        f"       tahun normal: {j(s.tahun_normal)} | abnormal (tidak dipakai): {j(s.tahun_abnormal)} | acuan laju pemulihan: {j(s.tahun_pemulihan)} dipakai {s.n_pemulihan} tahun",
        f"       tahun rujukan rata-rata: {j(s.tahun_rata)} | batas pencilan {s.outlier_cap:g}x | lantai pemulihan {'ya' if s.floor_pemulihan else 'tidak'}",
        "",
        "SHEET (satuan Rupiah penuh, kecuali dinyatakan lain):",
        f"1. Historis       : realisasi/APBD per akun, kolom {s.tahun_awal}…{s.tahun_dasar}. Sel yang dibutuhkan model tidak boleh kosong (tahun normal, pemulihan, rujukan, dasar, dan tahun sebelumnya).",
        "                    Tahun abnormal boleh kosong kecuali bila menjadi tahun sebelum tahun normal/pemulihan.",
        f"2. Dasar          : nilai tahun dasar ({s.tahun_dasar}). Baris trf_dau/trf_dbh/trf_dak OPSIONAL: isi bila ingin memodelkan pemotongan per jenis transfer pusat.",
        f"3. Pensiun        : jumlah pegawai pensiun per golongan untuk kolom {s.tahun_dasar}…{s.tahun_akhir} dan rata-rata gaji+tunjangan setahun (kolom gaji_tunjangan_per_tahun).",
        f"4. Jalur          : per tahun proyeksi ({s.tahun_dasar + 1}–{s.tahun_akhir}) jumlah ASN penerima TPP dan belanja gaji+tunjangan PPPK.",
        "5. Skenario_TPP   : total belanja TPP per tahun untuk tiap skenario kebijakan (tambah/hapus kolom). Skenario kebijakan lama dihitung otomatis.",
        "6. Kejadian       : (OPSIONAL) dampak langsung kejadian tak terduga. Kosongkan bila tidak ada.",
        "7. Asumsi_Kejadian: (OPSIONAL) perubahan asumsi model akibat kejadian per periode. Kosongkan bila tidak ada.",
        "8. Setup          : salinan setup sidebar (informasi; dipakai aplikasi untuk mendeteksi ketidakcocokan).",
        "",
        "KOLOM Kejadian: kejadian | tahun_mulai | tahun_selesai (kosong=1 tahun; 'akhir'=sampai akhir horizon) | komponen | mode (persen/rupiah) | nilai (negatif=berkurang) | keterangan",
        "KOMPONEN yang dikenali:",
    ] + [f"   {k:<22} {v}" for k, v in lab.items()] + [
        "PARAMETER Asumsi_Kejadian (mode ganti/tambah; nilai dalam persen):",
    ] + [f"   {k:<22} {v}" for k, v in ev.PARAMETER.items()] + [
        "",
        "Sheet Contoh_Kejadian hanya contoh (diabaikan aplikasi). Aplikasi juga menerima copy-paste langsung dari Excel ke tabel di tab 'Data & Asumsi'.",
    ]


def template_bytes(s: "pr.Setup", tables: dict | None = None) -> bytes:
    """Excel sesuai Setup. `tables` (opsional) = data yang sudah ada; diselaraskan ke setup (nilai yang cocok dipertahankan)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    tb = pr.align_tables(s, tables) if tables else pr.template_tables(s)
    K, A = ev.contoh_tables(s.proj_years, s.jenis_pemda)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"Petunjuk": _petunjuk(s)}).to_excel(xw, sheet_name="Petunjuk", index=False)
        s.to_df().to_excel(xw, sheet_name="Setup", index=False)
        for name in SHEETS:
            df = tb[name].copy()
            df.columns = [str(c) for c in df.columns]
            df.to_excel(xw, sheet_name=name, index=False)
        for name in EVENT_SHEETS:
            coerce_event_table(name, tb[name]).to_excel(xw, sheet_name=name, index=False)
        K.to_excel(xw, sheet_name="Contoh_Kejadian", index=False)
        A.to_excel(xw, sheet_name="Contoh_Kejadian", index=False, startrow=len(K) + 3)
    buf.seek(0)
    wb = load_workbook(buf)
    head = PatternFill("solid", fgColor="1F4E78")
    for ws in wb:
        if ws.title == "Petunjuk":
            ws.column_dimensions["A"].width = 150
            ws["A1"].font = Font(bold=True, size=12)
            continue
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for j, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
            name = str(col[0].value or "")
            ws.column_dimensions[get_column_letter(j)].width = 62 if name in ("uraian", "keterangan") else (34 if name in ("kejadian",) else 20)
        ws.freeze_panes = "C2" if ws.title in ("Historis", "Dasar") else "B2"
        if ws.title in ("Historis", "Dasar", "Pensiun", "Jalur", "Skenario_TPP"):
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, (int, float)) or c.value is None:
                        c.number_format = "#,##0.##"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
